"""Print CRRAH short names, names, and AR email domains as CSV."""
import csv
import sys
from openpyxl import load_workbook

def extract_crrah(path):
    """Return account holder entities from the workbook's CRRAH sheet."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["CRRAH"]
    entities = []

    for row in sheet.iter_rows(min_row=11, values_only=True):
        name, short_name, _, _, _, email = row[:6]
        if short_name:
            domain = email.split("@")[-1].lower() if email else ""
            entities.append({"short_name": short_name, "name": name, "domain": domain})

    workbook.close()
    return entities


if __name__ == "__main__":
    writer = csv.DictWriter(sys.stdout, fieldnames=["short_name", "name", "domain"])
    writer.writeheader()
    writer.writerows(extract_crrah(sys.argv[1]))
