"""Read the four supplied input files into ordinary SQLite columns."""
import calendar
import csv
from datetime import datetime
import io
import json
import zipfile

import openpyxl

HOLDINGS = "CRROwnershipofRecord10012026.zip"
ACCOUNTS = "List_of_Market_Participants_in_ERCOT_Region_20261011.xlsx"
HOURS = "CRR-Time-of-Use-Hours_2026-2029_UPD.xlsx"


def excel_rows(path, sheet, first_row):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        yield from workbook[sheet].iter_rows(min_row=first_row, values_only=True)
    finally:
        workbook.close()


def load_inputs(db, data_dir):
    db.execute("INSERT INTO reports VALUES (?, ?)", (1, "2026-10-01"))

    # Accounts workbook: CRRAH names and email domains.
    for row in excel_rows(data_dir / ACCOUNTS, "CRRAH", 11):
        if row[1] is not None:
            domain = row[5].rsplit("@", 1)[1].lower() if row[5] else None
            db.execute("INSERT INTO account_holders VALUES (?, ?, ?)",
                       (row[1], row[0], domain))

    # TOU workbook: published hours by month and block.
    for year, month, off_peak, peak_wd, peak_we, _total in excel_rows(
            data_dir / HOURS, "CRR Time of Use Hours", 3):
        if year is not None:
            delivery_month = f"{year:04d}-{list(calendar.month_name).index(month):02d}-01"
            db.executemany("INSERT INTO tou_hours VALUES (?, ?, ?)",
                           [(delivery_month, block, hours) for block, hours in
                            zip(("Off-peak", "PeakWD", "PeakWE"), (off_peak, peak_wd, peak_we))])

    # Holdings ZIP: contracts and source/sink placeholders.
    accounts = {row[0] for row in db.execute("SELECT account_holder FROM account_holders")}
    points = set()
    count, batch = 0, []
    insert = "INSERT INTO crr_holdings VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    with zipfile.ZipFile(data_dir / HOLDINGS) as archive, archive.open("Common_ExistingCRR.csv") as stream:
        reader = csv.DictReader(io.TextIOWrapper(stream, encoding="utf-8-sig"))
        reader.fieldnames = [name.strip() for name in reader.fieldnames]
        for row in reader:
            account, source, sink = row["AccountHolder"], row["Source"], row["Sink"]
            if account not in accounts:
                db.execute("INSERT INTO account_holders VALUES (?, NULL, NULL)", (account,))
                accounts.add(account)
            for point in (source, sink):
                if point not in points:
                    db.execute("INSERT INTO settlement_points (settlement_point) VALUES (?)", (point,))
                    points.add(point)
            start, end = (datetime.strptime(row[key], "%m/%d/%Y").date().isoformat()
                          for key in ("StartDate", "EndDate"))
            block = "Off-peak" if row["TimeOfUse"] == "OffPeak" else row["TimeOfUse"]
            batch.append((1, int(row["CRR_ID"]), int(row["SegmentID"]), account,
                          row["HedgeType"], row["CRRType"], source, sink, start, end,
                          block, float(row["MW"])))
            count += 1
            if len(batch) == 5000:
                db.executemany(insert, batch)
                batch.clear()
        db.executemany(insert, batch)

    # Topology JSON: locations and types for referenced settlement points.
    topology = json.loads((data_dir / "topology.json").read_text())
    for feature in topology["settlement_points"]["features"]:
        properties = feature["properties"]
        point_type = properties.get("sp_type")
        load_zone = properties.get("load_zone")
        geometry = feature.get("geometry")
        longitude, latitude = geometry["coordinates"] if geometry else (None, None)
        db.execute("""UPDATE settlement_points SET longitude=?, latitude=?, point_type=?,
                      load_zone=?, capacity_mw=? WHERE settlement_point=?""",
                   (longitude, latitude, point_type.upper() if point_type else None,
                    load_zone.upper() if load_zone else None,
                    properties.get("capacity_mw"), properties["sp_id"]))

    return count
